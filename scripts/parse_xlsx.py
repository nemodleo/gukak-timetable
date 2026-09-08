#!/usr/bin/env python3
"""
Parse every "2026년 N월(...).xlsx" instructor/room timetable in ./data into one
combined seed file: src/data/seed-2026.json  (Jan–Sep 2026, Aug file absent).

The generated file contains real student/teacher names, so it is git-ignored;
the copy committed to the repo is a small anonymised placeholder. Run this
script to regenerate the real one locally, or import it via /admin → 데이터.

Usage:
    python scripts/parse_xlsx.py [DATA_DIR] [OUT_JSON]
      DATA_DIR  default: ./data          (globs *.xlsx)
      OUT_JSON  default: src/data/seed-2026.json

Layout of each "Week N" sheet (data_only):
  - 7 day columns, each a block of 6 room columns  (I II III 大 中 우)
        day 0: M N O P Q R            (cols 13..18)   memo col S  (19)   time col L (12)
        day 1: X Y Z AA AB AC         (cols 24..29)   memo col AD (30)   time col W (23)
        day 2: AI AJ AK AL AM AN      (cols 35..40)   memo col AO (41)   time col AH(34)
        day 3: AT AU AV AW AX AY      (cols 46..51)   memo col AZ (52)   time col AS(45)
        day 4: BE BF BG BH BI BJ      (cols 57..62)   memo col BK (63)   time col BD(56)
        day 5: BP BQ BR BS BT BU      (cols 68..73)   memo col BV (74)   time col BO(67)
        day 6: CA CB CC CD CE CF      (cols 79..84)   memo col CG (85)   time col BZ(78)
  - lesson entries live in merged 6-row blocks starting at rows 8,14,20,26,32,38,44,50,56
        -> slot_index 0..8   (weekday 0..7 = 2h, 8 = 1h ; weekend 0..6, 6 = 1h)
  - the 7 grid dates come from the computed day headers K3/V3/AG3/AR3/BC3/BN3/BY3.

Every sheet has "Week 1..6"; some are stale copies (a repeated / near-empty week,
or "Week 6" left pointing at January). We de-duplicate per calendar day, keeping
whichever week sheet has the most filled cells for that day.

Each workbook is authoritative for its own month only (week sheets that spill into
a neighbouring month are dropped). The roster / target hours from each "SET UP"
sheet are emitted PER MONTH ("ym" = "2026-07"), not unioned — the roster changes
month to month. Gray-filled room cells (FFCCCCCC / FFD9D9D9) with no text are
emitted as kind:"block" (a non-teaching band).
"""
import json, re, sys, datetime, glob
from pathlib import Path
import openpyxl
from openpyxl.utils import column_index_from_string as ci

ROOT = Path(__file__).resolve().parent.parent
DATA_DIR = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "data"
OUT = Path(sys.argv[2]) if len(sys.argv) > 2 else ROOT / "src" / "data" / "seed-2026.json"

ROOMS = ["I", "II", "III", "大", "中", "우"]
DAY_BLOCK_START = ["M", "X", "AI", "AT", "BE", "BP", "CA"]
DAY_MEMO_COL = ["S", "AD", "AO", "AZ", "BK", "BV", "CG"]
DAY_HEADER_CELL = ["K3", "V3", "AG3", "AR3", "BC3", "BN3", "BY3"]
SLOT_ROWS = [8, 14, 20, 26, 32, 38, 44, 50, 56]

def _blocks(spec):
    return [{"start": s, "end": e} for s, e in spec]

WEEKDAY_SLOTS = _blocks([
    ("06:00", "08:00"), ("08:00", "10:00"), ("10:00", "12:00"), ("12:00", "14:00"),
    ("14:00", "16:00"), ("16:00", "18:00"), ("18:00", "20:00"), ("20:00", "22:00"),
    ("22:00", "23:00"),
])
WEEKEND_SLOTS = _blocks([
    ("05:00", "07:00"), ("07:00", "09:00"), ("09:00", "11:00"), ("11:00", "13:00"),
    ("13:00", "15:00"), ("15:00", "17:00"), ("17:00", "18:00"),
])

MONTHS = {m: i + 1 for i, m in enumerate(
    ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"])}


def file_month(path: Path) -> int | None:
    """month number from '2026년 7월.xlsx' / '2026-07.xlsx' / '... 3월(수정본).xlsx'."""
    m = re.search(r"(\d{1,2})\s*월", path.name)
    if m:
        return int(m.group(1))
    m = re.search(r"20\d{2}[-_](\d{1,2})", path.name)
    return int(m.group(1)) if m else None


def parse_header_date(text, target_year: int, target_month: int):
    """'Jul 01 (Wed)' -> date, choosing the year closest to the file's month."""
    m = re.match(r"([A-Za-z]{3})\s+(\d{1,2})", str(text or ""))
    if not m:
        return None
    mon, day = MONTHS[m.group(1)], int(m.group(2))
    tgt = target_year * 12 + (target_month - 1)
    best = None
    for yr in (target_year - 1, target_year, target_year + 1):
        try:
            d = datetime.date(yr, mon, day)
        except ValueError:
            continue
        dist = abs((yr * 12 + (mon - 1)) - tgt)
        if best is None or dist < best[0]:
            best = (dist, d)
    return best[1] if best else None


def clean(v):
    if v is None:
        return ""
    if isinstance(v, str):
        return v.replace("\r", "").strip()
    return str(v).strip()


def looks_like_time_label(s):
    return bool(re.search(r"(AM|PM)", s, re.I)) and "~" in s


# gray solid fills = 비수업 band (only meaningful when the cell has no text)
GRAY_FILLS = {"FFCCCCCC", "FFD9D9D9"}
# grade / accent fills -> our CellColor keys for a 직접 입력(text) cell
# (gray is NOT a text-cell color — a gray cell with text is just plain text)
FILL_COLOR = {
    "FFFFE599": "y", "FFEA9999": "r", "FFB6D7A8": "g", "FFA4C2F4": "b",
}


def _fill_rgb(cell):
    f = cell.fill
    if not f or f.patternType != "solid":
        return None
    rgb = getattr(f.fgColor, "rgb", None)
    return rgb.upper() if isinstance(rgb, str) else None


def is_gray(cell) -> bool:
    return _fill_rgb(cell) in GRAY_FILLS


def cell_color(cell):
    """CellColor key for a text cell's fill, or None (gray excluded)."""
    return FILL_COLOR.get(_fill_rgb(cell))


def ym_key(year: int, month: int) -> str:
    return f"{year}-{month:02d}"


def read_roster(su):
    out = []
    for r in range(11, 220):
        student = clean(su.cell(r, ci("D")).value)
        if not student:
            continue
        teacher = clean(su.cell(r, ci("F")).value)
        grade = su.cell(r, ci("E")).value
        target = su.cell(r, ci("G")).value
        grade_s = ""
        if grade is not None:
            try:
                grade_s = str(int(grade)) if float(grade).is_integer() else str(grade)
            except (TypeError, ValueError):
                grade_s = clean(grade)
        label = clean(su.cell(r, ci("C")).value) or f"{student} ({teacher}T)"
        out.append({
            "student_name": student,
            "grade": grade_s,
            "teacher_name": teacher,
            "target_hours": float(target) if isinstance(target, (int, float)) else None,
            "label": label,
        })
    return out


def main():
    files = sorted(
        (p for p in DATA_DIR.glob("*.xlsx") if not p.name.startswith("~$")),
        key=lambda p: (file_month(p) or 99, p.name),
    )
    if not files:
        sys.exit(f"no .xlsx files in {DATA_DIR}")

    year = 2026
    school_name = "국립국악고등학교"
    week_start = "mon"

    # per-month roster: ym -> ordered list of roster dicts (each tagged with ym)
    roster_by_ym: dict[str, list[dict]] = {}
    labels_by_ym: dict[str, set[str]] = {}

    # per calendar day, keep the richest version seen across sheets in its own file
    day_cells: dict[str, list[dict]] = {}
    day_memos: dict[str, list[str]] = {}
    day_score: dict[str, int] = {}

    for path in files:
        fm = file_month(path)
        if fm is None:
            print(f"  skip (no month in name): {path.name}")
            continue
        wb = openpyxl.load_workbook(path, data_only=True)
        ym = ym_key(year, fm)
        if "SET UP" in wb.sheetnames:
            su = wb["SET UP"]
            school_name = clean(su["D6"].value) or school_name
            year = int(su["D7"].value or year)
            ym = ym_key(year, fm)
            week_start = "mon" if "Monday" in clean(su["D8"].value) else "sun"
            roster = []
            for p in read_roster(su):
                roster.append({**p, "ym": ym})
            roster_by_ym[ym] = roster
            labels_by_ym[ym] = {p["label"] for p in roster}

        week_sheets = [s for s in wb.sheetnames if re.fullmatch(r"[Ww]eek ?\d+", s)]
        for wk in week_sheets:
            ws = wb[wk]
            dates = [parse_header_date(ws[c].value, year, fm) for c in DAY_HEADER_CELL]
            for di, d in enumerate(dates):
                if d is None:
                    continue
                # each monthly workbook is authoritative for its own month only.
                # week sheets spill into the previous/next month (e.g. Sep's
                # "Week 5" runs Sep 28–Oct 4); drop those out-of-month days so a
                # month with no source file stays empty.
                if d.year != year or d.month != fm:
                    continue
                iso = d.isoformat()
                block = ci(DAY_BLOCK_START[di])
                is_weekend = d.weekday() >= 5
                n_slots = len(WEEKEND_SLOTS) if is_weekend else len(WEEKDAY_SLOTS)

                cells, score = [], 0
                for si in range(n_slots):
                    row = SLOT_ROWS[si]
                    for room in ROOMS:
                        c = ws.cell(row, block + ROOMS.index(room))
                        val = clean(c.value)
                        if val and not looks_like_time_label(val):
                            score += 1
                            cells.append({"date": iso, "room": room, "slot_index": si,
                                          "value": val, "color": cell_color(c)})
                        elif is_gray(c):
                            # non-teaching band marked by a gray fill
                            cells.append({"date": iso, "room": room, "slot_index": si, "block": True})

                mcol = ci(DAY_MEMO_COL[di])
                memo_lines = []
                for row in range(4, 62):
                    mv = clean(ws.cell(row, mcol).value)
                    if mv and mv.lower() != "memo":
                        memo_lines.append(mv.replace("\n", " "))

                # keep the richest rendering of this day
                if score > day_score.get(iso, -1) or (
                    score == day_score.get(iso, -1) and memo_lines and not day_memos.get(iso)
                ):
                    day_score[iso] = score
                    day_cells[iso] = cells
                    day_memos[iso] = memo_lines

    pairings = []
    for ym in sorted(roster_by_ym):
        for i, p in enumerate(roster_by_ym[ym]):
            pairings.append({**p, "sort_order": i})

    cells_out = []
    for iso in sorted(day_cells):
        ymk = iso[:7]
        labels = labels_by_ym.get(ymk, set())
        for c in day_cells[iso]:
            base = {"date": iso, "room": c["room"], "slot_index": c["slot_index"]}
            if c.get("block"):
                # 회색 fill + 텍스트 없음 = 비수업(빈)
                cells_out.append({**base, "kind": "block", "label": None, "text": None, "color": None})
                continue
            v = c["value"]
            if v in labels:
                cells_out.append({**base, "kind": "pairing", "label": v, "text": None, "color": None})
            else:
                # 학생 배정이 아닌 자유 텍스트 = 비수업(메모) — 색 fill 있으면 유지
                cells_out.append({**base, "kind": "block", "label": None,
                                  "text": v.replace("\n", " "), "color": c.get("color")})

    memos_out = []
    for iso in sorted(day_memos):
        lines = []
        for m in day_memos[iso]:
            if m not in lines:
                lines.append(m)
        if not lines:
            continue
        memos_out.append({
            "date": iso,
            "lines": {
                f"{6 + (i * 30) // 60:02d}:{(i * 30) % 60:02d}": {"text": t}
                for i, t in enumerate(lines)
            },
        })

    out = {
        "settings": {
            "school_name": school_name,
            "year": year,
            "month": 9,
            "week_start": week_start,
            "rooms": ROOMS,
            "time_slots_weekday": WEEKDAY_SLOTS,
            "time_slots_weekend": WEEKEND_SLOTS,
        },
        "pairings": pairings,
        "cells": cells_out,
        "memos": memos_out,
    }
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(out, ensure_ascii=False, indent=2), encoding="utf-8")

    dd = sorted({c["date"] for c in cells_out})
    months = sorted({d[:7] for d in dd})
    print(f"wrote {OUT.relative_to(ROOT)}")
    print(f"  files:    {len(files)}  ({', '.join(p.name for p in files)})")
    ym_counts = {}
    for p in pairings:
        ym_counts[p["ym"]] = ym_counts.get(p["ym"], 0) + 1
    print(f"  pairings: {len(pairings)}  by month: "
          + ", ".join(f"{k}={v}" for k, v in sorted(ym_counts.items())))
    _blk = [c for c in cells_out if c["kind"] == "block"]
    print(f"  cells:    {len(cells_out)}  "
          f"(pairing={sum(1 for c in cells_out if c['kind']=='pairing')}, "
          f"block={len(_blk)}  [메모있음 {sum(1 for c in _blk if c['text'])}])")
    print(f"  memo days:{len(memos_out)}")
    print(f"  days:     {len(dd)}   span {dd[0]} .. {dd[-1]}")
    print(f"  months:   {', '.join(months)}")


if __name__ == "__main__":
    main()
